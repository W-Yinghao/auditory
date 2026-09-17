#!/usr/bin/env python3
import csv,hashlib,json,os,re,shutil
from collections import Counter,defaultdict
from pathlib import Path
import openpyxl
import matplotlib;matplotlib.use('Agg')
import matplotlib.pyplot as plt
R=Path(__file__).resolve().parents[1]; MAP=R/'private/inventory_001/file_path_map.csv'; O=R/'results/clinical_003'; P=R/'private/clinical_003'; C='file_e254c9a75c971ac7d4e76bf6df2b4030'
def norm(v): return re.sub(r'\s+',' ',str(v).strip()).lower() if v is not None else ''
def num(v):
 m=re.fullmatch(r'[-+]?\d+(?:\.\d+)?',norm(v).replace(',','')); return float(m.group()) if m else None
def source(fid):
 for r in csv.reader(MAP.open(encoding='utf-8')):
  if r and r[0]==fid:return Path(r[3])
 raise FileNotFoundError(fid)
def main():
 if not os.environ.get('SLURM_JOB_ID'):raise RuntimeError('Slurm required')
 O.mkdir(parents=True,exist_ok=True);P.mkdir(parents=True,exist_ok=True);p=source(C); h=hashlib.sha256(p.read_bytes()).hexdigest(); wb=openpyxl.load_workbook(p,read_only=False,data_only=True);ws=wb['Sheet1']; rows=list(ws.iter_rows(values_only=True));hdr=list(rows[0]);
 def merged(row,col):
  for mr in ws.merged_cells.ranges:
   if mr.min_row<=row<=mr.max_row and mr.min_col<=col<=mr.max_col:return ws.cell(mr.min_row,mr.min_col).value
  return None
 def ci(label):return next(j for j,v in enumerate(hdr) if norm(v)==norm(label))
 ni=ci('姓名'); ai=ci('实际月龄'); di=ci('HA使用时长（月）'); mi=ci('MUSS得分（%）'); ii=ci('IT-MAIS/MAIS得分（%）'); si=ci('SIR得分'); cap=ci('CAP得分'); wi=ci('每日佩戴时长')
 rec=[]; suffix_date=0;suffix_device=0
 for rn in range(2,ws.max_row+1):
  r=list(rows[rn-1]);
  if not any(v is not None and norm(v) for v in r):continue
  raw=str(r[ni]) if r[ni] is not None else ''; m=re.match(r'^([\u4e00-\u9fff]+)',raw); base=m.group(1) if m else ''; suffix=raw[len(base):] if base else raw
  if re.search(r'\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}(?!\d)',suffix):suffix_date+=1
  if re.search(r'(?i)(ci|ha|双模式|bimodal|助听器|耳蜗)',suffix):suffix_device+=1
  g=merged(rn,1) or r[0]; rec.append({'rn':rn,'raw_name':raw,'suffix':suffix,'candidate_key':hashlib.sha256(base.encode()).hexdigest()[:16] if base else '', 'group':g,'age':num(r[ai]),'dur':num(r[di]),'muss':num(r[mi]),'itmais':num(r[ii]),'sir':num(r[si]),'cap':num(r[cap]),'wear':r[wi]})
 groups=Counter(norm(r['group']) for r in rec); cl=Counter(r['candidate_key'] for r in rec if r['candidate_key']); histbad=[]; ha60=Counter(); scales=defaultdict(list)
 for r in rec:
  if norm(r['group'])!='nh':
   if r['dur'] is not None and 0<=r['dur']<=60: ha60['all_0_60']+=1
   if r['dur'] is not None and r['dur']<0:ha60['negative']+=1
   for k in ('cap','sir','itmais','muss'):scales['ha_'+k].append(r[k])
  else:
   for k in ('cap','sir','itmais','muss'):scales['nh_'+k].append(r[k])
  exp='nh' if norm(r['group'])=='nh' else ('0m' if r['dur']==0 else ('<6m' if r['dur'] is not None and 0<r['dur']<6 else ('6-11m' if r['dur'] is not None and 6<=r['dur']<12 else ('12-24m' if r['dur'] is not None and 12<=r['dur']<=24 else ('25-36m' if r['dur'] is not None and 24<r['dur']<=36 else ('37-60m' if r['dur'] is not None and 36<r['dur']<=60 else ('>60m' if r['dur'] is not None and r['dur']>60 else 'unknown')))))))
  if norm(r['group']) not in (exp,'nh') and exp!='unknown':histbad.append({'source_row':r['rn'],'group':r['group'],'duration_stratum':exp})
 result={'job_id':os.environ['SLURM_JOB_ID'],'source_file_id':C,'source_sha256':h,'records':len(rec),'candidate_clusters':len(cl),'candidate_cluster_record_counts':dict(sorted(cl.items())),'group_record_counts':dict(groups),'HA_records':sum(norm(r['group'])!='nh' for r in rec),'NH_records':sum(norm(r['group'])=='nh' for r in rec),'HA_0_60_continuous_record_count':ha60.get('all_0_60',0),'negative_duration_count':ha60.get('negative',0),'historical_group_duration_disagreement_rows':histbad,'name_suffix_date_like_count':suffix_date,'name_suffix_device_like_count':suffix_device,'suffix_interpretation':'Date/device-like suffixes are label evidence only, not confirmed visit dates or device state.','scale_by_group':{k:{'n':len([x for x in v if x is not None]),'observed_max':max([x for x in v if x is not None],default=None),'observed_max_9_count':sum(x==9 for x in v if x is not None),'observed_max_5_count':sum(x==5 for x in v if x is not None),'observed_max_100_count':sum(x==100 for x in v if x is not None)} for k,v in scales.items()}}
 (O/'clinical_003_summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
 with (P/'clinical_rows_clean.csv').open('w',newline='',encoding='utf-8') as f:
  w=csv.writer(f);w.writerow(['source_row','candidate_key','group','raw_name','suffix','age_months','duration_months','MUSS','IT_MAIS_MAIS','SIR','CAP','daily_wear']);[w.writerow([r[k] for k in ('rn','candidate_key','group','raw_name','suffix','age','dur','muss','itmais','sir','cap','wear')]) for r in rec]
 (P/'candidate_identity_map.csv').write_text('candidate_key,status\n'+'\n'.join(f'{k},candidate_only' for k in sorted(cl))+'\n',encoding='utf-8')
 # Keep 002 provenance-rich copy private and remove absolute paths from its public summary.
 old=R/'results/clinical_002/clinical_002_summary.json'
 if old.exists():
  d=json.loads(old.read_text(encoding='utf-8')); (P/'clinical_002_summary_with_paths.json').write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
  for b in d.get('workbooks',[]):b.pop('absolute_path_used',None)
  old.write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
 # Separate percentage and ordinal panels; all plots are record-level.
 fig,ax=plt.subplots(1,2,figsize=(9,4));
 for k in ('ha_muss','ha_itmais','nh_muss','nh_itmais'):
  v=[x for x in scales[k] if x is not None]
  if v:ax[0].hist(v,bins=12,alpha=.45,label=k.replace('_',' '))
 for k in ('ha_cap','ha_sir','nh_cap','nh_sir'):
  v=[x for x in scales[k] if x is not None]
  if v:ax[1].hist(v,bins=10,alpha=.45,label=k.replace('_',' '))
 ax[0].set(xlabel='Percentage scale, record level',ylabel='Records');ax[1].set(xlabel='Ordinal scale, record level',ylabel='Records');ax[0].legend(fontsize=7);ax[1].legend(fontsize=7);fig.tight_layout();fig.savefig(O/'clinical_003_scales_by_group.png',dpi=140);plt.close(fig)
if __name__=='__main__':main()
