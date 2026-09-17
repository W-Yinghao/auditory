#!/usr/bin/env python3
import csv,hashlib,json,re,os
from collections import Counter,defaultdict
from pathlib import Path
import openpyxl
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
R=Path(__file__).resolve().parents[1]; MAP=R/'private/inventory_001/file_path_map.csv'; S=R/'private/probe_001'; O=R/'results/clinical_002'; P=R/'private/clinical_002'
IDS=['file_e254c9a75c971ac7d4e76bf6df2b4030','file_8a9c2767bddf88cb1bc2531e0dd4ea00','file_d4bd0c2924e725928f61c2756853cb38','file_225d0f1d5180f1dc47fb96cedbde60d7']
def norm(v): return re.sub(r'\s+',' ',str(v).strip()).lower() if v is not None else ''
def val(v):
 m=re.fullmatch(r'[-+]?\d+(?:\.\d+)?',norm(v).replace(',','')); return float(m.group()) if m else None
def path_for(fid):
 for r in csv.reader(MAP.open(encoding='utf-8')):
  if r and r[0]==fid:return Path(r[3])
 raise FileNotFoundError(fid)
def main():
 if not os.environ.get('SLURM_JOB_ID'): raise RuntimeError('audit must run under Slurm')
 O.mkdir(parents=True,exist_ok=True);P.mkdir(parents=True,exist_ok=True); result={'job_id':os.environ['SLURM_JOB_ID'],'primary':{},'workbooks':[]}
 paths={i:path_for(i) for i in IDS}
 for fid,p in paths.items():
  wb=openpyxl.load_workbook(p,read_only=False,data_only=True); sheets=[]
  for ws in wb.worksheets:
   rows=list(ws.iter_rows(values_only=True)); non=[r for r in rows[1:] if any(v is not None and norm(v) for v in r)]; names=sum(bool(next((v for v in r if isinstance(v,str) and norm(v)),'')) for r in non)
   sheets.append({'title':ws.title,'rows_including_header':len(rows),'nonempty_rows':len(non),'merged_ranges':[str(x) for x in ws.merged_cells.ranges],'nonempty_first_text_fields':names})
  result['workbooks'].append({'file_id':fid,'absolute_path_used':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':sheets})
 # Primary workbook: fill group only through confirmed merged ranges; preserve column numbers.
 p=paths[IDS[0]]; wb=openpyxl.load_workbook(p,read_only=False,data_only=True); ws=wb['Sheet1']; rows=list(ws.iter_rows(values_only=True)); hdr=list(rows[0]); merges=[str(x) for x in ws.merged_cells.ranges]
 def merged_value(row,col):
  for mr in ws.merged_cells.ranges:
   if mr.min_row<=row<=mr.max_row and mr.min_col<=col<=mr.max_col: return ws.cell(mr.min_row,mr.min_col).value
  return None
 rec=[]
 for rn in range(2,ws.max_row+1):
  vals=list(next(ws.iter_rows(min_row=rn,max_row=rn,values_only=True)))
  if not any(v is not None and norm(v) for v in vals):continue
  q={f'col_{j+1}':vals[j] if j<len(vals) else None for j in range(len(hdr))}; g=merged_value(rn,1) or (vals[0] if vals else None); q['group_filled_from_merged_range']=g; nm=norm(q.get('col_2')); q['candidate_key']=hashlib.sha256(re.sub(r'([_-]?\d+|[（(][^）)]*[）)])$','',nm).encode()).hexdigest()[:16] if nm else ''; rec.append(q)
 def colnum(label): return next((j+1 for j,h in enumerate(hdr) if norm(h)==norm(label)),None)
 def nums(label):
  j=colnum(label); return [val(r[f'col_{j}']) for r in rec if j and val(r[f'col_{j}']) is not None]
 age=nums('实际月龄'); dur=nums('HA使用时长（月）'); scale={k:nums(k) for k in ('CAP得分','SIR得分','IT-MAIS/MAIS得分（%）','MUSS得分（%）')}; groups=Counter(norm(r['group_filled_from_merged_range']) for r in rec)
 result['primary']={'file_id':IDS[0],'sheet':'Sheet1','records':len(rec),'candidate_clusters':len({r['candidate_key'] for r in rec if r['candidate_key']}),'group_record_counts':dict(groups),'merged_ranges':merges,'age_range':[min(age),max(age)] if age else None,'duration_range':[min(dur),max(dur)] if dur else None,'missing_by_column':{f'col_{j+1}':sum(r[f'col_{j+1}'] is None for r in rec) for j in range(len(hdr))},'scale_ranges':{k:[min(v),max(v)] if v else None for k,v in scale.items()},'observed_max_9_count':sum(v==9 for v in scale['CAP得分']),'MUSS_87_count':sum(v==87 for v in scale['MUSS得分（%）']),'daily_wear_nonempty':sum(bool(norm(r[f'col_{colnum("每日佩戴时长")}'])) for r in rec),'source_hash':hashlib.sha256(p.read_bytes()).hexdigest()}
 with (P/'clinical_rows_clean.csv').open('w',newline='',encoding='utf-8') as f:
  w=csv.writer(f);w.writerow(['candidate_key','group','source_row']+[f'col_{i+1}' for i in range(len(hdr))]);[w.writerow([r['candidate_key'],r['group_filled_from_merged_range'],i+2]+[r[f'col_{j+1}'] for j in range(len(hdr))]) for i,r in enumerate(rec)]
 (P/'candidate_identity_map.csv').write_text('candidate_key,status\n'+'\n'.join(f"{k},candidate_only" for k in sorted({r['candidate_key'] for r in rec if r['candidate_key']}))+'\n',encoding='utf-8');(O/'clinical_002_summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
 pairs=[(val(r[f'col_{next(j+1 for j,h in enumerate(hdr) if norm(h)==norm("HA使用时长（月）"))}']),val(r[f'col_{next(j+1 for j,h in enumerate(hdr) if norm(h)==norm("实际月龄"))}'])) for r in rec if val(r[f'col_{next(j+1 for j,h in enumerate(hdr) if norm(h)==norm("HA使用时长（月）"))}']) is not None and val(r[f'col_{next(j+1 for j,h in enumerate(hdr) if norm(h)==norm("实际月龄"))}']) is not None]
 fig,axs=plt.subplots(1,2,figsize=(10,4)); axs[0].scatter([a for a,b in pairs],[b for a,b in pairs],s=10);axs[0].set(xlabel='HA duration (months), record level',ylabel='Age (months), record level');
 for k,v in scale.items():
  if v: axs[1].hist(v,bins=12,alpha=.45,label=k.split('（')[0])
 axs[1].set_xlabel('Score, record level');axs[1].legend(fontsize=7);fig.tight_layout();fig.savefig(O/'clinical_002_distributions.png',dpi=140);plt.close(fig)
if __name__=='__main__':main()
