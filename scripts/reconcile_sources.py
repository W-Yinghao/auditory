"""Explicit clinical/vendor/BDF candidate links, with no assumed concurrent labels."""
import csv,hashlib,json,os,re,traceback
from collections import Counter,defaultdict
from datetime import datetime
from pathlib import Path
import openpyxl
from audit_mff import dump,table,j
BASE=Path(__file__).resolve().parents[1]
assert os.environ.get('SLURM_JOB_ID')
out=BASE/'results/linkage_001';out.mkdir(exist_ok=False);priv=BASE/'private/linkage_001';priv.mkdir(mode=0o700,exist_ok=False)
mapping=list(csv.DictReader((BASE/'private/inventory_001/file_path_map.csv').open()));pathbyid={r['file_id']:Path(r['absolute_path']) for r in mapping};idbypath={str(v):k for k,v in pathbyid.items()}
def namekey(x):
 s=str(x or '').strip();m=re.match(r'[\u4e00-\u9fff]+',s)
 return m.group() if m else re.sub(r'\s+',' ',s).lower()
def dated(x):
 m=re.search(r'(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})',str(x))
 if not m:return ''
 try:return datetime(*map(int,m.groups())).date().isoformat()
 except ValueError:return ''
def tid(prefix,x):return prefix+hashlib.sha256(str(x).encode()).hexdigest()[:12]
cid='file_e254c9a75c971ac7d4e76bf6df2b4030';wb=openpyxl.load_workbook(pathbyid[cid],read_only=False,data_only=True);ws=wb['Sheet1'];clinical=[]
for rn in range(2,ws.max_row+1):
 vals=[c.value for c in ws[rn]]
 if not vals[1]:continue
 name=vals[1];clinical.append({'clinical_row_id':f'C{rn:04d}','source_row':rn,'name':str(name),'namekey':namekey(name),'label_date':dated(name),'age_months':vals[3],'device_use_months':vals[4]})
vendors=[];sources=[];errors=[];erp=[]
for r in mapping:
 p=Path(r['absolute_path'])
 if p.name not in ('recordInformation.json','erp.json'):continue
 try:
  d=json.loads(p.read_text());fid=r['file_id'];gid=tid('B',p.parent)
  if not isinstance(d,dict):erp.append({'file_id':fid,'status':'null_or_nonobject_export'});continue
  if p.name=='recordInformation.json':
   vendors.append({'file_id':fid,'candidate_acquisition_id':gid,'namekey':namekey(d.get('PatientName')),'PatientName':d.get('PatientName'),'PatientGUID':d.get('PatientGUID'),'PatientID':d.get('PatientID'),'BirthDate':d.get('BirthDate'),'ExamTime':d.get('ExamTime'),'StartRecordTime':d.get('StartRecordTime'),'RecordTime':d.get('RecordTime'),'ExamGUID':d.get('ExamGUID'),'date':dated(d.get('StartRecordTime')) or dated(d.get('ExamTime')),'folder':str(p.parent)})
   sources.append({'file_id':fid,'candidate_acquisition_id':gid,'device_name':d.get('DeviceName'),'sample_rate':d.get('SampleRate'),'record_montage':d.get('RecordMontageName'),'data_file_info_count':len(d.get('DataFileInformations') or []),'patient_guid_present':bool(d.get('PatientGUID')),'patient_id_present':bool(d.get('PatientID')),'birth_date_present':bool(d.get('BirthDate')),'exam_time_present':bool(d.get('ExamTime')),'start_record_time_present':bool(d.get('StartRecordTime')),'device_state':'unknown'})
  else:
   settings=d.get('OverlapSettings') or {};averages=d.get('EpochAverages') or [];mmn=d.get('MMNEpoch')
   erp.append({'file_id':fid,'candidate_acquisition_id':gid,'status':'parsed','n_condition_averages':len(averages),'MMNS1':d.get('MMNS1'),'MMNS2':d.get('MMNS2'),'MMN_dictionary_present':isinstance(mmn,dict),'MMN_n_channels':len(mmn.get('DataChannel',[])) if isinstance(mmn,dict) else 0,'low_cut':settings.get('LowCut'),'high_cut':settings.get('HighCut'),'reference':settings.get('Reference'),'epoch_left_s':settings.get('EpochLeft'),'epoch_right_s':settings.get('EpochRight'),'data_level':'evoked','source_trials_available_in_json':bool(d.get('SourceDatas'))})
 except Exception as e:errors.append({'file_id':r['file_id'],'error':repr(e)})
vbygroup={v['candidate_acquisition_id']:v for v in vendors};namegroups=defaultdict(list)
for v in vendors:namegroups[v['namekey']].append(v)
links=[];participants=[];byclinical=defaultdict(list)
for c in clinical:byclinical[c['namekey']].append(c)
for nk,vs in namegroups.items():
 pid=tid('P',nk);guids=set(v['PatientGUID'] for v in vs if v['PatientGUID']);dob=set(v['BirthDate'] for v in vs if v['BirthDate'])
 participants.append({'participant_id':pid,'identity_status':'candidate_name_group','n_vendor_acquisitions':len(vs),'n_patient_guids':len(guids),'n_birthdate_strings':len(dob),'n_clinical_rows':len(byclinical[nk]),'cross_cohort_status':'not_resolved'})
 for v in vs:
  matches=byclinical[nk];date_matches=[c for c in matches if c['label_date'] and c['label_date']==v['date']]
  selected=date_matches if date_matches else matches
  if not selected:links.append({'candidate_acquisition_id':v['candidate_acquisition_id'],'participant_id':pid,'clinical_row_id':'','clinical_match_status':'none','evidence':'no_name_candidate'})
  for c in selected:
   links.append({'candidate_acquisition_id':v['candidate_acquisition_id'],'participant_id':pid,'clinical_row_id':c['clinical_row_id'],'clinical_match_status':'candidate','evidence':'name_and_label_date' if c in date_matches else 'name_only','n_candidate_clinical_rows':len(selected),'concurrent_scale_date_status':'unknown','eeg_clinical_gap_days':''})
bdf=list(csv.DictReader((BASE/'results/signal_001/bdf_pairing.csv').open()));bh=json.loads((BASE/'private/other_eeg_001/bdf_private_headers.json').read_text());bh={r['file_id']:r for r in bh}
recs=[];files=[];deltas=[]
for b in bdf:
 v=vbygroup.get(b['candidate_acquisition_id']);a=bh[b['signal_file_id']];e=bh[b['event_file_id']]
 def time(h):return datetime.strptime(h['date']+' '+h['time'],'%d.%m.%y %H.%M.%S')
 delta=(time(e)-time(a)).total_seconds();deltas.append(delta)
 rec={'recording_id':b['candidate_acquisition_id'],'recording_status':'paired_export_candidate_acquisition','participant_id':tid('P',v['namekey']) if v else '',
 'participant_status':'candidate_name_group' if v else 'unresolved','visit_id':'','visit_status':'unresolved','device_state':'unknown','vendor_metadata_present':v is not None,'event_header_minus_signal_header_s':delta,'pairing_status':'header_start_equal' if delta==0 else 'timing_reconciliation_required','data_level':'raw_export','qc_status':'headers_events_checked'}
 if v:rec['vendor_date_matches_bdf']=v['date']==time(a).date().isoformat()
 recs.append(rec)
 for fid,role in [(b['signal_file_id'],'signal'),(b['event_file_id'],'events')]:files.append({'file_id':fid,'recording_id':rec['recording_id'],'file_role':role,'data_version_id':'source_BDF_export','parent_file_id':'','pairing_status':rec['pairing_status']})
 for r in mapping:
  if str(Path(r['absolute_path']).parent)==a['folder'] and Path(r['absolute_path']).suffix.lower() in ('.set','.fdt','.json'):
   files.append({'file_id':r['file_id'],'recording_id':rec['recording_id'],'file_role':Path(r['absolute_path']).suffix[1:],'data_version_id':'same_directory_derivative_or_metadata','parent_file_id':'','pairing_status':'candidate_directory_only'})
for name,rows in [('participant_index',participants),('recording_index',recs),('files_to_recordings',files),('clinical_link',links),('vendor_acquisition_metadata',sources),('vendor_evoked_metadata',erp)]:table(out/(name+'.csv'),rows)
dump(priv/'clinical_rows.json',clinical);dump(priv/'vendor_identity_records.json',vendors);dump(priv/'errors.json',errors)
guids=defaultdict(set)
for v in vendors:
 if v['PatientGUID']:guids[v['PatientGUID']].add(v['namekey'])
summary={'job_id':os.environ['SLURM_JOB_ID'],'clinical_records':len(clinical),'clinical_name_candidates':len(byclinical),'vendor_records':len(vendors),'vendor_name_candidates':len(namegroups),'vendor_unique_patient_guids':len(guids),'patient_guids_with_multiple_name_keys':sum(len(x)>1 for x in guids.values()),'name_candidates_with_multiple_patient_guids':sum(r['n_patient_guids']>1 for r in participants),'name_candidates_with_multiple_dob_strings':sum(r['n_birthdate_strings']>1 for r in participants),'clinical_rows_with_date_in_label':sum(bool(c['label_date']) for c in clinical),'candidate_link_rows':len(links),'candidate_link_evidence':dict(Counter(r['evidence'] for r in links)),'acquisitions_with_exactly_one_candidate_clinical_row':sum(sum(l['candidate_acquisition_id']==r['recording_id'] and bool(l['clinical_row_id']) for l in links)==1 for r in recs),'event_minus_signal_header_seconds':dict(Counter(deltas)),'vendor_date_bdf_matches':sum(r.get('vendor_date_matches_bdf',False) for r in recs),'vendor_erp_status':dict(Counter(r['status'] for r in erp)),'confirmed_concurrent_eeg_scale_links':0,'errors':len(errors),'notice':'Name groups and dates embedded in row labels are candidate evidence, not verified concurrent clinical visits.'}
dump(out/'summary.json',summary);print(j(summary))
