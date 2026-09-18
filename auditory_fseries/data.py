"""Build the frozen, archival F-series data package.

This module consumes the already reconstructed Phase 1 epoch packages.  It
does not open raw EEG files or alter filtering/reference choices.  Clinical
values are read from the registered HA workbook so that the literal source
headers remain part of the private lineage record.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
from pathlib import Path
from typing import Any, Iterable, Mapping

import numpy as np


CANONICAL_CHANNELS = (
    "Fp1", "Fp2", "Fz", "F3", "F4", "F7", "F8", "Cz", "C3", "C4",
    "T3", "T4", "Pz", "P3", "P4", "T5", "T6", "Oz", "O1", "O2",
)
FREQUENCIES = (500, 1000, 2000, 4000)
WORKBOOK_FILE_ID = "file_e254c9a75c971ac7d4e76bf6df2b4030"
EPSILON = 1e-12


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value)).casefold()


def _number(value: Any) -> float | None:
    if isinstance(value, (int, float, np.integer, np.floating)) and not isinstance(value, bool):
        result = float(value)
    else:
        raw = _text(value).replace(",", "").rstrip("%")
        if not raw:
            return None
        try:
            result = float(raw)
        except ValueError:
            return None
    return result if math.isfinite(result) else None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _as_path(root: Path, config: Mapping[str, Any], key: str, default: Path) -> Path:
    value = config.get(key)
    if value is None and isinstance(config.get("paths"), Mapping):
        value = config["paths"].get(key)
    if value is None:
        return default
    path = Path(value)
    return path if path.is_absolute() else root / path


def _uniform_indices(n: int, maximum: int) -> np.ndarray:
    if n < 1:
        raise ValueError("no accepted epochs")
    k = min(int(maximum), n)
    if k == n:
        return np.arange(n, dtype=np.int64)
    # Endpoints and chronological order are fixed.  This depends only on the
    # accepted-trial count, never on code or a clinical value.
    indices = np.rint(np.linspace(0, n - 1, k)).astype(np.int64)
    if len(np.unique(indices)) != k:
        raise RuntimeError("uniform trial selection produced duplicate indices")
    return indices


def _finite_or_error(name: str, values: np.ndarray) -> None:
    if not np.isfinite(values).all():
        raise ValueError(f"nonfinite {name}")


def _record_statistics(data: np.ndarray, selected: np.ndarray) -> np.ndarray:
    """Return 60 deterministic per-channel Hjorth summaries."""
    x = np.asarray(data[selected], dtype=np.float64)
    if x.ndim != 3 or x.shape[1:] != (20, 176):
        raise ValueError(f"selected epoch data must have shape (n,20,176), got {x.shape}")
    _finite_or_error("selected epoch data", x)
    x = x - np.mean(x, axis=-1, keepdims=True)
    var_x = np.maximum(np.var(x, axis=-1), EPSILON)
    d1 = np.diff(x, axis=-1)
    d2 = np.diff(d1, axis=-1)
    var_d1 = np.maximum(np.var(d1, axis=-1), EPSILON)
    var_d2 = np.maximum(np.var(d2, axis=-1), EPSILON)
    mobility = np.sqrt(var_d1 / var_x) * 250.0
    complexity = np.sqrt(var_d2 / var_d1) / np.sqrt(var_d1 / var_x)
    components = (
        np.log(np.maximum(var_x, EPSILON)),
        np.log(np.maximum(mobility, EPSILON)),
        np.log(np.maximum(complexity, EPSILON)),
    )
    features = np.concatenate([np.median(component, axis=0) for component in components])
    _finite_or_error("Hjorth feature", features)
    if features.shape != (60,):
        raise RuntimeError(f"unexpected feature shape {features.shape}")
    return features


def record_features(
    data: np.ndarray,
    times: np.ndarray,
    accepted: np.ndarray,
    codes: np.ndarray,
    samples: np.ndarray,
    config: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Compute one record's Z features and quality values.

    The accepted mask is applied before chronological uniform sampling.  The
    optional ``times`` argument is checked for the frozen 176-sample epoch
    shape but is otherwise not used, preserving the supplied epoch window.
    """
    cfg = dict(config or {})
    max_epochs = int(cfg.get("maximum_epochs", cfg.get("max_epochs", 256)))
    if max_epochs < 1:
        raise ValueError("maximum_epochs must be positive")
    data = np.asarray(data)
    accepted = np.asarray(accepted)
    codes = np.asarray(codes)
    samples = np.asarray(samples)
    ordinal = np.asarray(cfg.get("ordinal", np.arange(len(data), dtype=np.int64)))
    times = np.asarray(times)
    if data.ndim != 3 or data.shape[1:] != (20, 176):
        raise ValueError(f"epoch data must have shape (n,20,176), got {data.shape}")
    n = data.shape[0]
    if accepted.ndim != 2 or accepted.shape[0] != n or accepted.shape[1] < 1:
        raise ValueError("accepted mask shape mismatch")
    if codes.shape[0] != n or samples.shape[0] != n or ordinal.shape[0] != n:
        raise ValueError("code/sample shape mismatch")
    if times.size and times.shape != (176,):
        raise ValueError("epoch times must have shape (176,)")
    primary = accepted[:, 0].astype(bool) & np.isin(codes, (1, 2))
    indices = np.flatnonzero(primary)
    if len(indices) < 64:
        raise ValueError(f"fewer than 64 accepted code1/2 epochs: {len(indices)}")
    if not np.isfinite(samples[indices]).all():
        raise ValueError("nonfinite accepted sample coordinates")
    # Stable tie break on the original ordinal/index keeps duplicate sample
    # positions deterministic without using labels.
    order = np.lexsort((ordinal[indices], samples[indices]))
    chronological = indices[order]
    selected = chronological[_uniform_indices(len(chronological), max_epochs)]
    z = _record_statistics(data, selected)
    accepted_data = np.asarray(data[indices], dtype=np.float64)
    _finite_or_error("accepted epoch data", accepted_data)
    ptp = np.ptp(accepted_data, axis=-1).max(axis=-1)
    _finite_or_error("accepted scalp peak-to-peak", ptp)
    duration_s = _number(cfg.get("duration_s"))
    if duration_s is None or duration_s < 0:
        raise ValueError("finite nonnegative duration_s required")
    q = np.asarray(
        [
            np.log1p(len(indices)),
            1.0 - len(indices) / float(n),
            np.log(max(float(np.median(ptp)), EPSILON)),
            np.log(max(duration_s, EPSILON)),
        ],
        dtype=np.float64,
    )
    _finite_or_error("quality features", q)
    return {
        "Z": z,
        "Q": q,
        "accepted_count": int(len(indices)),
        "stored_count": int(n),
        "selected_indices": selected.astype(np.int64),
        "chronological_indices": chronological.astype(np.int64),
    }


def _workbook_path(root: Path, config: Mapping[str, Any], inventory_path: Path) -> Path:
    direct = config.get("ha_workbook_path")
    if direct is None and isinstance(config.get("paths"), Mapping):
        direct = config["paths"].get("ha_workbook_path")
    if direct is not None:
        path = Path(direct)
        return path if path.is_absolute() else root / path
    file_id = str(config.get("ha_workbook_file_id", WORKBOOK_FILE_ID))
    for row in _read_csv(inventory_path):
        if _text(row.get("file_id")) == file_id:
            path = Path(_text(row.get("absolute_path")))
            if not path.exists():
                raise FileNotFoundError(path)
            return path
    raise KeyError(f"registered HA workbook file id not found: {file_id}")


def _read_workbook(path: Path) -> tuple[dict[int, dict[str, Any]], dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment contract
        raise RuntimeError("openpyxl is required to read the registered HA workbook") from exc
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    if "Sheet1" not in workbook.sheetnames:
        raise KeyError("registered HA workbook lacks Sheet1")
    sheet = workbook["Sheet1"]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not rows:
        raise ValueError("empty HA workbook")
    headers = [_text(value) for value in rows[0]]
    normalized = [_norm(value) for value in headers]

    def column(label: str) -> int:
        target = _norm(label)
        if target not in normalized:
            raise KeyError(f"verified HA workbook column missing: {label}")
        return normalized.index(target)

    cols = {
        "name": column("姓名"),
        "group": 0,
        "age_months": column("实际月龄"),
        "duration_months": column("HA使用时长（月）"),
        "MUSS": column("MUSS得分（%）"),
        "IT_MAIS_MAIS": column("IT-MAIS/MAIS得分（%）"),
    }
    threshold_cols: dict[str, list[int]] = {}
    for panel, device in (("unaided", "裸耳"), ("aided", "助听")):
        for side, side_label in (("right", "右"), ("left", "左")):
            key = f"{side}_{panel}"
            threshold_cols[key] = [
                column(f"{side_label}耳{frequency}Hz({device})") for frequency in FREQUENCIES
            ]
    merged_groups: dict[int, Any] = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_col <= 1 <= merged.max_col:
            value = sheet.cell(merged.min_row, 1).value
            for row_number in range(merged.min_row, merged.max_row + 1):
                merged_groups[row_number] = value
    out: dict[int, dict[str, Any]] = {}
    sequence = 0
    for source_row, values in enumerate(rows[1:], 2):
        if not any(_text(value) for value in values):
            continue
        sequence += 1
        raw = {key: [values[index] if index < len(values) else None for index in indexes]
               for key, indexes in threshold_cols.items()}
        pta: dict[str, float | None] = {}
        for key, values_for_side in raw.items():
            parsed = [_number(value) for value in values_for_side]
            pta[f"{key}_pta"] = (
                float(sum(parsed) / 4.0) if len(parsed) == 4 and all(value is not None for value in parsed) else None
            )
        unaided = [pta["right_unaided_pta"], pta["left_unaided_pta"]]
        aided = [pta["right_aided_pta"], pta["left_aided_pta"]]
        out[source_row] = {
            "source_row": source_row,
            "sequence": sequence,
            "name": _text(values[cols["name"]]) if cols["name"] < len(values) else "",
            "group": _text(merged_groups.get(source_row, values[cols["group"]] if cols["group"] < len(values) else "")),
            "age_months": _number(values[cols["age_months"]]),
            "duration_months": _number(values[cols["duration_months"]]),
            "MUSS": _number(values[cols["MUSS"]]),
            "IT_MAIS_MAIS": _number(values[cols["IT_MAIS_MAIS"]]),
            "better_unaided_pta": min(unaided) if all(value is not None for value in unaided) else None,
            "better_aided_pta": min(aided) if all(value is not None for value in aided) else None,
            "raw_thresholds": {key: [_text(value) for value in values_for_side] for key, values_for_side in raw.items()},
            **pta,
        }
    evidence = {
        "sheet": "Sheet1",
        "literal_headers": {key: headers[index] for key, index in cols.items()},
        "threshold_headers": {
            key: [headers[index] for index in indexes] for key, indexes in threshold_cols.items()
        },
        "rows": len(out),
    }
    return out, evidence


def _is_nh(value: Any) -> bool:
    normalized = _norm(value)
    return normalized in {"nh", "健听", "健聽", "normal", "control", "对照", "對照"}


def _source_row_from_clinical_id(value: Any) -> int:
    match = re.fullmatch(r"C(\d{4})", _text(value))
    if match is None:
        raise ValueError(f"clinical row id must be C####: {value!r}")
    return int(match.group(1))


def _bool(value: Any) -> bool:
    return _norm(value) in {"true", "1", "yes", "y"}


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                fields.append(key)
                seen.add(key)
    if not fields:
        fields = ["status"]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _source_path_map(path: Path) -> dict[str, str]:
    return {_text(row.get("file_id")): _text(row.get("absolute_path")) for row in _read_csv(path)}


def prepare(root: Path | str, private: Path | str, public: Path | str, config: Mapping[str, Any] | Path | str) -> dict[str, Any]:
    """Prepare the frozen exploratory cohort and write its private arrays."""
    if not os.environ.get("SLURM_JOB_ID"):
        raise RuntimeError("F-series data preparation must run under Slurm")
    root = Path(root)
    private = Path(private)
    public = Path(public)
    if isinstance(config, (Path, str)):
        config_path = Path(config)
        if not config_path.is_absolute():
            config_path = root / config_path
        config_obj = json.loads(config_path.read_text(encoding="utf-8"))
    else:
        config_path = None
        config_obj = dict(config)
    # The orchestration runner creates and snapshots the run directory before
    # calling this adapter.  Direct callers may supply an empty directory.
    private.mkdir(parents=True, exist_ok=True, mode=0o700)
    public.mkdir(parents=True, exist_ok=True)
    os.chmod(private, 0o700)

    linked_path = _as_path(root, config_obj, "linked_index", root / "private/phase2_cohort_001/linked_index.csv")
    clinical_path = _as_path(root, config_obj, "clinical_rows", root / "private/clinical_003/clinical_rows_clean.csv")
    inventory_path = _as_path(root, config_obj, "inventory", root / "private/inventory_001/file_path_map.csv")
    source_manifest_path = _as_path(root, config_obj, "source_manifest", root / "results/phase1_sources_001/source_manifest.csv")
    epochs_root = _as_path(root, config_obj, "epochs_root", root / "results/phase1_epochs_001")
    workbook_path = _workbook_path(root, config_obj, inventory_path)
    paths = [linked_path, clinical_path, inventory_path, source_manifest_path, workbook_path]
    if config_path is not None:
        paths.append(config_path)
    for path in paths:
        if not path.exists():
            raise FileNotFoundError(path)
    inputs = {
        "job_id": os.environ["SLURM_JOB_ID"],
        "paths": {
            str(path): {"sha256": _sha256(path), "size_bytes": path.stat().st_size} for path in paths
        },
    }
    (private / "inputs.json").write_text(json.dumps(inputs, ensure_ascii=False, indent=2), encoding="utf-8")

    linked = _read_csv(linked_path)
    clinical_rows = _read_csv(clinical_path)
    clinical_by_source = {_text(row.get("source_row")): row for row in clinical_rows}
    source_rows = _read_csv(source_manifest_path)
    source_by_recording = {_text(row.get("recording_id")): row for row in source_rows}
    workbook_rows, workbook_evidence = _read_workbook(workbook_path)
    (private / "workbook_schema.json").write_text(json.dumps(workbook_evidence, ensure_ascii=False, indent=2), encoding="utf-8")

    minimum = int(config_obj.get("population", {}).get("minimum_primary_accepted_epochs_total", 64))
    maximum = int(config_obj.get("features", {}).get("maximum_epochs", 256))
    accepted_rows: list[dict[str, Any]] = []
    flow: dict[str, int] = {"linked_rows": len(linked), "earliest_ha_rows": 0, "metadata_eligible_rows": 0, "feature_rows": 0}
    exclusions: dict[str, int] = {}
    features_z: list[np.ndarray] = []
    features_q: list[np.ndarray] = []
    targets_a: list[float] = []
    targets_v: list[float] = []
    groups: list[str] = []
    recordings: list[str] = []

    def exclude(reason: str) -> None:
        exclusions[reason] = exclusions.get(reason, 0) + 1

    for link in linked:
        if _text(link.get("cohort")) != "HA" or not _bool(link.get("eligible_measurement_identity_index")):
            continue
        flow["earliest_ha_rows"] += 1
        if _bool(link.get("identity_dob_conflict")):
            exclude("dob_conflict")
            continue
        source = source_by_recording.get(_text(link.get("recording_id")))
        if source is None or _text(source.get("cohort_label")) != "HA" or _text(source.get("source_gate")) != "eligible_technical_measurement":
            exclude("source_not_eligible_ha")
            continue
        if not (_bool(link.get("strong_unique_link")) and _text(link.get("clinical_link_evidence")) == "name_and_label_date" and _text(link.get("clinical_link_count")) == "1"):
            exclude("link_not_unique_name_labeldate")
            continue
        clinical_id = _text(link.get("clinical_row_id"))
        clinical = None
        # clinical_rows_clean is keyed by source_row; the link carries C####.
        source_row_text = _text(link.get("clinical_row_id"))
        try:
            source_row_number = _source_row_from_clinical_id(source_row_text)
        except ValueError:
            exclude("clinical_row_id_invalid")
            continue
        source_row_text = str(source_row_number)
        clinical = clinical_by_source.get(source_row_text)
        worksheet = workbook_rows.get(int(source_row_text))
        if clinical is None or worksheet is None:
            exclude("clinical_source_row_missing")
            continue
        if _text(clinical.get("raw_name")) != _text(worksheet.get("name")):
            raise ValueError("registered clinical row does not match original workbook name")
        for old_key, new_key in [('MUSS', 'MUSS'), ('IT_MAIS_MAIS', 'IT_MAIS_MAIS'),
                                 ('age_months', 'age_months'), ('duration_months', 'duration_months')]:
            if _number(clinical.get(old_key)) != worksheet.get(new_key):
                raise ValueError("registered clinical values do not match original workbook row")
        if _is_nh(clinical.get("group")) or _is_nh(worksheet.get("group")):
            exclude("clinical_row_not_HA")
            continue
        flow["metadata_eligible_rows"] += 1
        age = worksheet.get("age_months")
        duration = worksheet.get("duration_months")
        target_a = worksheet.get("IT_MAIS_MAIS")
        target_v = worksheet.get("MUSS")
        if any(value is None for value in (age, duration, target_a, target_v)):
            exclude("required_clinical_value_missing")
            continue
        if age < 0 or duration < 0:
            exclude("negative_age_or_duration")
            continue
        if not (0.0 <= target_a <= 100.0 and 0.0 <= target_v <= 100.0):
            exclude("target_out_of_range")
            continue
        recording_id = _text(link.get("recording_id"))
        epoch_path = epochs_root / recording_id / "epochs.npz"
        if not epoch_path.exists():
            exclude("epoch_package_missing")
            continue
        with np.load(epoch_path, allow_pickle=False) as archive:
            required = {"data_uv", "times_s", "accepted", "codes", "samples_0based", "event_indices_1based", "channels"}
            if not required.issubset(archive.files):
                raise ValueError(f"{epoch_path}: missing required epoch arrays")
            data = archive["data_uv"]
            channels = tuple(_text(value) for value in archive["channels"].tolist())
            if channels != CANONICAL_CHANNELS:
                raise ValueError(f"{epoch_path}: noncanonical channel order")
            accepted = archive["accepted"]
            primary = accepted[:, 0].astype(bool) & np.isin(archive["codes"], (1, 2))
            if int(primary.sum()) < minimum:
                exclude("fewer_than_64_accepted_epochs")
                continue
            expected_times = np.linspace(-0.2, 0.5, 176, dtype=np.float64)
            if archive["times_s"].shape != (176,) or not np.allclose(archive["times_s"], expected_times, atol=1e-12, rtol=0.0):
                raise ValueError(f"{epoch_path}: epoch times do not match frozen -0.2..0.5 s grid")
            stats = record_features(
                data, archive["times_s"], accepted, archive["codes"], archive["samples_0based"],
                {"maximum_epochs": maximum, "duration_s": source.get("duration_s"), "ordinal": archive["event_indices_1based"]},
            )
        flow["feature_rows"] += 1
        row: dict[str, Any] = {
            "participant_id": _text(link.get("participant_id")),
            "recording_id": recording_id,
            "clinical_row_id": clinical_id,
            "worksheet_source_row": int(source_row_text),
            "worksheet_name": worksheet.get("name", ""),
            "worksheet_group": worksheet.get("group", ""),
            "cohort": _text(link.get("cohort")),
            "source_gate": _text(link.get("source_gate")),
            "ordinal": _text(link.get("ordinal")),
            "vendor_exam_time": _text(link.get("vendor_exam_time")),
            "source_duration_s": _number(source.get("duration_s")),
            "stored_epochs": stats["stored_count"],
            "accepted_epochs": stats["accepted_count"],
            "selected_epochs": len(stats["selected_indices"]),
            "age_months": age,
            "ha_duration_months": duration,
            "MUSS_literal_percent": target_v,
            "IT_MAIS_MAIS_literal_percent": target_a,
            "better_unaided_pta_source_units": worksheet["better_unaided_pta"],
            "better_aided_pta_source_units": worksheet["better_aided_pta"],
            "r_unaided_pta": worksheet["right_unaided_pta"],
            "l_unaided_pta": worksheet["left_unaided_pta"],
            "r_aided_pta": worksheet["right_aided_pta"],
            "l_aided_pta": worksheet["left_aided_pta"],
            "raw_thresholds_json": json.dumps(worksheet["raw_thresholds"], ensure_ascii=False, sort_keys=True),
        }
        for key, values in worksheet["raw_thresholds"].items():
            row[f"{key}_raw_json"] = json.dumps(values, ensure_ascii=False)
        accepted_rows.append(row)
        features_z.append(stats["Z"])
        features_q.append(stats["Q"])
        targets_a.append(float(target_a))
        targets_v.append(float(target_v))
        groups.append(_text(link.get("participant_id")))
        recordings.append(recording_id)

    epoch_paths = {
        str(epochs_root / row["recording_id"] / "epochs.npz"): {
            "sha256": _sha256(epochs_root / row["recording_id"] / "epochs.npz"),
            "size_bytes": (epochs_root / row["recording_id"] / "epochs.npz").stat().st_size,
        }
        for row in accepted_rows
    }
    inputs["epoch_packages"] = epoch_paths
    (private / "inputs.json").write_text(json.dumps(inputs, ensure_ascii=False, indent=2), encoding="utf-8")

    _write_csv(private / "cohort.csv", accepted_rows)
    n = len(accepted_rows)
    arrays = {
        "C": np.asarray([[row["age_months"], np.log1p(row["ha_duration_months"]), row["better_unaided_pta_source_units"] if row["better_unaided_pta_source_units"] is not None else np.nan, row["better_aided_pta_source_units"] if row["better_aided_pta_source_units"] is not None else np.nan] for row in accepted_rows], dtype=np.float64).reshape(n, 4),
        "Q": np.asarray(features_q, dtype=np.float64).reshape(n, 4),
        "Z": np.asarray(features_z, dtype=np.float64).reshape(n, 60),
        "A": np.asarray(targets_a, dtype=np.float64),
        "V": np.asarray(targets_v, dtype=np.float64),
        "groups": np.asarray(groups, dtype="U"),
        "recordings": np.asarray(recordings, dtype="U"),
    }
    for key, value in arrays.items():
        if key in {"Q", "Z", "A", "V"}:
            _finite_or_error(key, value)
    if not np.isfinite(arrays['C'][:, :2]).all() or np.isinf(arrays['C']).any():
        raise ValueError("invalid clinical features beyond permitted missing PTA")
    if len(set(groups)) != n or any(not group for group in groups):
        raise ValueError("one nonempty candidate identity per record required")
    np.savez(private / "data.npz", **arrays)
    n_groups = len(set(groups))
    min_groups = int(config_obj.get("population", {}).get("minimum_identity_groups", 30))
    status = "READY" if n_groups >= min_groups else "SUPPORT_LIMITED"
    summary = {
        "status": status,
        "flow": flow,
        "counts": {"records": n, "identity_groups": n_groups, "eligible_linked_rows": flow["metadata_eligible_rows"]},
        "exclusions": exclusions,
        "feature_dimensions": {"C": [n, 4], "Q": [n, 4], "Z": [n, 60]},
        "source_schema": workbook_evidence,
    }
    (private / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary
